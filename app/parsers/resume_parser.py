"""Extract plain text from PDF and DOCX resumes. No LLM calls."""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

_PDF_MAGIC = b"%PDF"
_ZIP_MAGIC = b"PK"
_BULLET_CHARS = ("•", "●", "○", "▪", "◦", "‣")
_NUMBERED_RE = re.compile(r"^(\d+|[A-Za-z])[.)]\s+")


class ResumeParseError(Exception):
    """Base error for resume document extraction."""


class UnsupportedResumeFormatError(ResumeParseError):
    """Raised when the file is not a supported PDF or DOCX resume."""


class EmptyResumeDocumentError(ResumeParseError):
    """Raised when the document has no extractable text."""


class CorruptedResumeDocumentError(ResumeParseError):
    """Raised when the document cannot be opened or decoded."""


def extract_resume_text(
    source: bytes | Path | str,
    *,
    filename: str | None = None,
) -> str:
    """Return clean resume text from a PDF or DOCX file.

    The Resume Agent can call this for any supported file. Raw bytes are never
    sent to an LLM from this module.
    """
    data, name = _read_source(source, filename)
    if not data:
        raise EmptyResumeDocumentError("Resume file is empty")

    kind = _detect_kind(data, name)
    try:
        if kind == "pdf":
            text = _extract_pdf(data)
        elif kind == "docx":
            text = _extract_docx(data)
        else:
            raise UnsupportedResumeFormatError(
                "Unsupported resume format. Provide a PDF or DOCX file."
            )
    except ResumeParseError:
        raise
    except Exception as exc:
        raise CorruptedResumeDocumentError(
            f"Resume document could not be read: {exc}"
        ) from exc

    text = _normalize_extracted_text(text)
    if not text:
        raise EmptyResumeDocumentError(
            "Resume document contains no extractable text"
        )
    return text


def _read_source(
    source: bytes | Path | str, filename: str | None
) -> tuple[bytes, str | None]:
    if isinstance(source, (bytes, bytearray, memoryview)):
        return bytes(source), filename
    path = Path(source)
    try:
        data = path.read_bytes()
    except FileNotFoundError as exc:
        raise ResumeParseError(f"Resume file not found: {path}") from exc
    except OSError as exc:
        raise CorruptedResumeDocumentError(
            f"Resume file could not be read: {exc}"
        ) from exc
    return data, filename or path.name


def _detect_kind(data: bytes, filename: str | None) -> str | None:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".docx":
        return "docx"
    if suffix in {".doc", ".rtf", ".txt", ".png", ".jpg", ".jpeg"}:
        return None
    if data.startswith(_PDF_MAGIC):
        return "pdf"
    if _looks_like_docx(data):
        return "docx"
    return None


def _looks_like_docx(data: bytes) -> bool:
    if not data.startswith(_ZIP_MAGIC):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
    except zipfile.BadZipFile:
        return False
    return "word/document.xml" in names


def _extract_pdf(data: bytes) -> str:
    import fitz

    try:
        document = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise CorruptedResumeDocumentError(
            "PDF resume could not be opened"
        ) from exc

    try:
        if document.is_encrypted:
            unlocked = document.authenticate("")
            if not unlocked:
                raise CorruptedResumeDocumentError(
                    "PDF resume is encrypted and cannot be read"
                )
        if document.page_count == 0:
            raise EmptyResumeDocumentError("PDF resume has no pages")
        pages: list[str] = []
        for page in document:
            pages.append(_pdf_page_text(page))
        return "\n\n".join(part for part in pages if part)
    finally:
        document.close()


def _pdf_page_text(page: object) -> str:
    payload = page.get_text("dict")
    sizes: list[float] = []
    for block in payload.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                sizes.append(float(span.get("size") or 0))
    median = sorted(sizes)[len(sizes) // 2] if sizes else 11.0
    heading_min = max(median + 1.0, median * 1.12)

    lines: list[str] = []
    for block in payload.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = line.get("spans") or []
            raw = "".join(str(span.get("text") or "") for span in spans).strip()
            if not raw:
                continue
            max_size = max((float(span.get("size") or 0) for span in spans), default=0)
            flags = 0
            for span in spans:
                flags |= int(span.get("flags") or 0)
            is_bold = bool(flags & 16)
            is_heading = max_size >= heading_min or (
                is_bold and max_size >= median and len(raw) <= 80
            )
            formatted = _format_line(raw, heading=is_heading)
            if is_heading and lines and lines[-1] != "":
                lines.append("")
            lines.append(formatted)
            if is_heading:
                lines.append("")
    return "\n".join(lines)


def _extract_docx(data: bytes) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(io.BytesIO(data))
    except (PackageNotFoundError, ValueError, KeyError, OSError) as exc:
        raise CorruptedResumeDocumentError(
            "DOCX resume could not be opened"
        ) from exc

    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            if lines and lines[-1] != "":
                lines.append("")
            continue
        style_name = ""
        if paragraph.style is not None and paragraph.style.name:
            style_name = paragraph.style.name.lower()
        if style_name.startswith("heading") or style_name == "title":
            if lines and lines[-1] != "":
                lines.append("")
            lines.append(text)
            lines.append("")
        elif _is_list_paragraph(paragraph):
            lines.append(f"- {_strip_bullet_prefix(text)}")
        else:
            lines.append(_format_line(text, heading=False))

    for table in document.tables:
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            cells = [cell for cell in cells if cell]
            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines)


def _is_list_paragraph(paragraph: object) -> bool:
    style = getattr(paragraph, "style", None)
    style_name = (style.name or "").lower() if style is not None else ""
    if "list" in style_name:
        return True
    p_pr = getattr(getattr(paragraph, "_p", None), "pPr", None)
    if p_pr is not None and getattr(p_pr, "numPr", None) is not None:
        return True
    return False


def _format_line(text: str, *, heading: bool) -> str:
    if heading:
        return text
    if _NUMBERED_RE.match(text):
        return f"- {_NUMBERED_RE.sub('', text).strip()}"
    stripped = _strip_bullet_prefix(text)
    if stripped != text:
        return f"- {stripped}"
    return text


def _strip_bullet_prefix(text: str) -> str:
    value = text.lstrip()
    for marker in _BULLET_CHARS:
        if value.startswith(marker):
            return value[len(marker) :].strip()
    if value[:2] in {"- ", "* ", "– ", "— "}:
        return value[2:].strip()
    return text.strip()


def _normalize_extracted_text(text: str) -> str:
    cleaned = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).rstrip() for line in cleaned.split("\n")]
    collapsed: list[str] = []
    for line in lines:
        if line == "" and (not collapsed or collapsed[-1] == ""):
            continue
        collapsed.append(line)
    return "\n".join(collapsed).strip()
