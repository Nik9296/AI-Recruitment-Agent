"""Generate a Word document that explains this project end-to-end."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    output = root / "AI-Recruitment-Agent-Step-by-Step-Guide.docx"

    doc = Document()
    doc.add_heading("AI Recruitment Agent - Full Project Guide", level=0)
    doc.add_paragraph(
        f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}."
    )
    doc.add_paragraph(
        "This document explains the complete architecture, workflow, APIs, database "
        "flow, and how to run the project step by step."
    )

    doc.add_heading("1) What this project is", level=1)
    doc.add_paragraph(
        "AI Recruitment Agent is a recruiter-assist platform. It analyzes job "
        "descriptions, sources candidates from internal ATS data, matches and ranks "
        "them, supports recruiter review decisions, creates outreach drafts, and "
        "helps schedule interviews."
    )
    add_bullets(
        doc,
        [
            "Backend: FastAPI + Pydantic + SQLAlchemy + Alembic",
            "Frontend: Next.js + React recruiter dashboard",
            "AI: Gemini via LangChain for JD and semantic tasks",
            "Orchestration: LangGraph pipeline for sourcing -> matching -> ranking",
            "Database: PostgreSQL",
        ],
    )

    doc.add_heading("2) Core design principles", level=1)
    add_bullets(
        doc,
        [
            "No invented candidate data; all candidate records come from stored ATS data.",
            "No autonomous hiring decisions.",
            "No outreach without recruiter approval.",
            "No automatic calendar booking without explicit confirmation step.",
            "Secrets are environment-driven (no hardcoded credentials).",
        ],
    )

    doc.add_heading("3) High-level architecture", level=1)
    doc.add_paragraph("A) Frontend (Next.js dashboard)")
    add_bullets(
        doc,
        [
            "Login with recruiter credentials.",
            "Jobs list and job shortlist pages.",
            "Candidate profile page with review actions.",
            "Outreach draft preview/send controls.",
            "Interview invite/propose/confirm controls.",
        ],
    )
    doc.add_paragraph("B) Backend (FastAPI)")
    add_bullets(
        doc,
        [
            "Versioned APIs under /api/v1.",
            "Health and readiness endpoints: /health and /ready.",
            "Route groups: jobs, recruiter, outreach, pipeline, scheduling.",
        ],
    )
    doc.add_paragraph("C) Data layer")
    add_bullets(
        doc,
        [
            "SQLAlchemy models for jobs, candidates, matches, workflow transitions, outreach, email sends, interviews.",
            "Alembic migrations manage schema evolution.",
            "Repositories isolate database operations from business logic.",
        ],
    )

    doc.add_heading("4) Step-by-step business workflow", level=1)
    doc.add_paragraph("Step 1 - Job Description (JD) analysis")
    add_bullets(
        doc,
        [
            "Recruiter submits raw JD text.",
            "JD Agent extracts structured JobProfile (title, skills, location, experience, etc.).",
            "Endpoint: POST /api/v1/jobs/analyze",
        ],
    )

    doc.add_paragraph("Step 2 - Job storage")
    add_bullets(
        doc,
        [
            "Structured job profile is saved in PostgreSQL.",
            "Endpoint: POST /api/v1/jobs",
        ],
    )

    doc.add_paragraph("Step 3 - Candidate sourcing (internal ATS only)")
    add_bullets(
        doc,
        [
            "Sourcing service converts job requirements into search queries.",
            "Repository fetches existing candidates by skills/title patterns.",
            "No external scraping and no invented candidates.",
            "Endpoint: POST /api/v1/recruiter/jobs/{job_id}/source",
        ],
    )

    doc.add_paragraph("Step 4 - Matching")
    add_bullets(
        doc,
        [
            "Matching agent computes deterministic scores (required skills, preferred skills, experience, title, etc.).",
            "Optional semantic embedding score is included when configured.",
            "Output is structured MatchResult entries.",
        ],
    )

    doc.add_paragraph("Step 5 - Ranking")
    add_bullets(
        doc,
        [
            "Ranking engine orders candidates by weighted score.",
            "Top candidates become shortlist records for recruiter review.",
        ],
    )

    doc.add_paragraph("Step 6 - Pipeline orchestration")
    add_bullets(
        doc,
        [
            "LangGraph flow: source -> match_and_rank -> record_shortlist -> human_gate.",
            "Pipeline explicitly ends before outreach/hiring automation.",
            "Endpoint: POST /api/v1/recruiter/jobs/{job_id}/pipeline/run",
        ],
    )

    doc.add_paragraph("Step 7 - Recruiter review decisions")
    add_bullets(
        doc,
        [
            "Recruiter can approve, reject, or hold each candidate.",
            "Review actions map to workflow states with transition validation.",
            "Endpoints: /approve, /reject, /hold under recruiter routes.",
        ],
    )

    doc.add_paragraph("Step 8 - Outreach handoff and drafts")
    add_bullets(
        doc,
        [
            "Approved + eligible open-to-work candidate can be handed to outreach queue.",
            "Outreach agent generates draft messages (not auto-sent).",
            "Email send endpoint enforces approval/open-to-work gates and duplicate prevention.",
        ],
    )

    doc.add_paragraph("Step 9 - Interview scheduling")
    add_bullets(
        doc,
        [
            "Candidate can be invited after outreach/contact state.",
            "Scheduling agent proposes slots.",
            "Recruiter confirms one slot.",
            "Calendar can run in dry-run or live Google mode (if configured).",
        ],
    )

    doc.add_paragraph("Step 10 - Final human decisions")
    add_bullets(
        doc,
        [
            "System supports final workflow states but does not autonomously hire.",
            "Human recruiter remains in control.",
        ],
    )

    doc.add_heading("5) Workflow states and transitions", level=1)
    add_bullets(
        doc,
        [
            "NEW -> MATCHED -> SHORTLISTED -> PENDING_REVIEW",
            "PENDING_REVIEW -> APPROVED_FOR_OUTREACH / REJECTED / ON_HOLD",
            "APPROVED_FOR_OUTREACH -> CONTACTED (system gate)",
            "CONTACTED / RESPONDED -> INTERVIEW_INVITED -> INTERVIEW_SCHEDULED",
            "HIRED and WITHDRAWN are terminal business outcomes with controlled transitions.",
        ],
    )

    doc.add_heading("6) Important API groups", level=1)
    add_bullets(
        doc,
        [
            "Jobs APIs: analyze and persist job profiles.",
            "Recruiter APIs: login, dashboard views, candidate decisions, workflow history.",
            "Pipeline APIs: source and end-to-end run.",
            "Outreach APIs: generate/get drafts and controlled send.",
            "Scheduling APIs: invite, propose, confirm, and retrieve interview plans.",
            "Health APIs: /health, /ready.",
        ],
    )

    doc.add_heading("7) Database flow (simplified)", level=1)
    add_bullets(
        doc,
        [
            "jobs table stores parsed job profile.",
            "candidates table stores candidate profile/resume metadata.",
            "job_candidates table stores match/ranking relation and recruiter status.",
            "workflow transition tables keep auditable state changes.",
            "outreach/email tables track draft/send state.",
            "interviews table stores proposed and confirmed interview slots.",
        ],
    )

    doc.add_heading("8) Security and configuration model", level=1)
    add_bullets(
        doc,
        [
            "All secrets come from .env/environment variables.",
            "Recruiter auth uses token-based verification.",
            "Production mode validates required config at startup.",
            "Readiness endpoint reports status without leaking secrets.",
        ],
    )

    doc.add_heading("9) How to run on another device (recommended)", level=1)
    doc.add_paragraph("Method A - Docker (easiest)")
    add_bullets(
        doc,
        [
            "Install Docker Desktop.",
            "Copy/clone project.",
            "Run: python run.py",
            "Open: http://localhost:3000",
            "Login: recruiter / recruiter-review (or values in .env).",
            "Stop: python stop.py",
        ],
    )

    doc.add_paragraph("Method B - Local Python + Node + PostgreSQL")
    add_bullets(
        doc,
        [
            "Install Python 3.11+, Node 22+, PostgreSQL.",
            "Create environment: python scripts/bootstrap_env.py",
            "Create venv and install requirements.",
            "Run migrations: python -m alembic upgrade head",
            "Seed demo data: python scripts/seed_recruiter_dashboard.py",
            "Run API and frontend in separate terminals.",
        ],
    )

    doc.add_heading("10) Typical end-to-end user journey", level=1)
    add_bullets(
        doc,
        [
            "Recruiter logs in and opens a job.",
            "Runs pipeline to source + rank candidates.",
            "Reviews candidate evidence and sets approve/reject/hold.",
            "Generates outreach draft and sends only with explicit approval.",
            "Invites candidate to interview and confirms a proposed slot.",
        ],
    )

    doc.add_heading("11) Current limitations and expected behavior", level=1)
    add_bullets(
        doc,
        [
            "External sourcing (LinkedIn/job boards) is intentionally not active.",
            "Email/calendar are dry-run until live credentials are configured.",
            "Semantic embedding components require LLM/API configuration.",
            "Workflow blocks invalid transitions with clear conflict responses.",
        ],
    )

    doc.add_heading("12) Quick troubleshooting checklist", level=1)
    add_bullets(
        doc,
        [
            "If login fails: verify recruiter username/password in .env.",
            "If dashboard is blank: check API /health and /ready endpoints.",
            "If DB errors appear: verify DATABASE_URL and migration status.",
            "If send/invite fails: check workflow status constraints.",
            "If AI analysis fails: verify GOOGLE_API_KEY and model settings.",
        ],
    )

    doc.add_paragraph("")
    doc.add_paragraph(
        "End of guide. This file is generated by scripts/generate_project_word_guide.py."
    )

    doc.save(output)
    print(f"Created: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
